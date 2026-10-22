from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile

root=Path(__file__).resolve().parents[2]
source=root/'output/docx/image-agent-batch-guide-20261010-21tasks.docx'
out=root/'output/docx/image-agent-batch-guide-20261013-21tasks-hard.docx'
doc=Document(source)
names=[('VocaloidEnsemble','中文V家六位歌手的经典版本合影','2.0'),('GenshinEnsemble','原神指定官方衣装与地区归属','3.0'),('ParisTrocaderoView','1900年巴黎世博会时期的反向实景','3.0')]

# Remove only the three v1 world-knowledge pages, retaining all old tasks/media.
first=next(p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('19 '))
node=first._p
while node is not None and node.tag!=qn('w:sectPr'):
 following=node.getnext();node.getparent().remove(node);node=following

for p in doc.paragraphs:
 if p.text=='2026-10-10 · 21题派发版':p.text='2026-10-13 · 21题难题派发版'
 elif p.text.startswith('21 个题组'):
  p.text='21 个题组。前18题沿用原指南；第19题采用中文V家v2，第20题采用原神指定衣装v3，第21题采用1900年巴黎实景v3。三题分别测试角色与衣装版本、地区归属，以及历史年代与空间关系。循环精灵图、连续场景和旧渡船漫画不纳入本批次。'
 elif p.text.startswith('第19至21题是'):
  p.text='第19至21题是无参考图的世界知识测试，分别采用v2、v3、v3题面。整个生成过程禁止外部搜索；仅提交本题提示词原文，不追加角色外观答案、地区分组答案、历史照片、地图或评审资料。每个模型配置每题生成一次，保留首轮缺陷，不自动重试。旧版结果不可作为新版结果或输入。'
 elif p.text.startswith('V家同屏、原神同屏和巴黎实景。'):
  p.text='难题包含中文V家经典衣装、原神指定官方衣装与国家分组、1900年巴黎反向机位。三题均禁止搜索且无参考图。知识正确性与画面质量分开评审；单次结果只用于初步观察，不能宣布模型上限或稳定排名。'
 for number,(tid,title,version) in enumerate(names,19):
  if tid in p.text and p.text.startswith(str(number)):
   p.text=f'{number:02d}  {title}  |  {tid}  v{version}'

for number,(tid,title,version) in enumerate(names,19):
 p=doc.add_paragraph(f'{number:02d}  {title}','Heading 1');p.paragraph_format.page_break_before=True
 p=doc.add_paragraph(f'任务 ID：{tid} | 版本：{version} | 分类：世界知识 文生图')
 for r in p.runs:r.font.size=Pt(9)
 doc.add_paragraph('输入与执行条件','Heading 2')
 doc.add_paragraph('无参考图。使用下列提示词原文，不附评审材料或此前生成图。禁止调用外部搜索、检索、浏览器或素材库。')
 doc.add_paragraph('提示词原文','Heading 2')
 prompt=(root/'Image'/tid/'_versions'/('v'+version.split('.')[0])/'PROJECT_PROMPT.md').read_text(encoding='utf-8').strip()
 for block in prompt.split('\n\n'):
  doc.add_paragraph(block)
 doc.add_paragraph('交付与记录','Heading 2')
 doc.add_paragraph(f'保存原始成图、完整实际提示词、模型与设置、耗时和输出SHA-256；meta.json须记录task_version为{version}。记录实际输出尺寸、外部工具轨迹及缺失信息；生成工具未公开的内部行为记为未验证。不同题面版本独立入库和评审。')

for section in doc.sections:
 for p in section.footer.paragraphs:
  for r in p.runs:r.text=r.text.replace('2026-10-10','2026-10-13')
for s in doc.styles:
 for b in s.element.xpath('.//w:pBdr'):b.getparent().remove(b)
for b in doc.element.xpath('//w:pBdr'):b.getparent().remove(b)
doc.core_properties.subject='21题难题派发版 世界知识v2 v3 v3 原始参考图内嵌 不包含标准答案和生成样例'
doc.save(out)
assert len(doc.inline_shapes)==13
assert sum('任务 ID：' in p.text for p in doc.paragraphs)==21
with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
 for n in a.namelist():
  if n.startswith('word/media/'):assert a.read(n)==b.read(n)
print(f'{out}: 21 tasks, 13 reference placements, world knowledge versions 2/3/3')
