"""Create an editable companion from the PDF snapshot, not live task files."""
from pathlib import Path
from io import BytesIO
import re
import pdfplumber
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
source = root/'output/pdf/image-agent-batch-guide-20261007.pdf'
out = root/'output/docx/image-agent-batch-guide-20261007.docx'
out.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.8)
for name, size in [('Normal',11),('Title',25),('Heading 1',18),('Heading 2',13),('Caption',9)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.1
doc.styles['Normal'].paragraph_format.widow_control = True
footer = section.footer.paragraphs[0]
footer.text = '图像任务库  Agent 批量执行指南  2026-10-07    '
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.style = 'Caption'
reader = PdfReader(source)
all_text = []
with pdfplumber.open(source) as pdf:
    for index, page in enumerate(pdf.pages):
        first_paragraph = True
        lines = [l for l in page.extract_text_lines() if l['bottom'] < page.height-42]
        events = []
        for line in lines:
            chars = line.pop('chars')
            line['size'] = max(c['size'] for c in chars)
            line['bold'] = any('Bold' in c['fontname'] for c in chars)
            events.append((line['top'],'text',line))
        # Each original PDF page uses distinct image XObjects in placement order.
        images = {image.name.rsplit('.',1)[0]:image for image in reader.pages[index].images}
        locations = sorted(page.images, key=lambda im: im['top'])
        for location in locations:
            image=images[location['name']]
            events.append((location['top'],'image',(location,image.data)))
        events.sort(key=lambda e:e[0])
        pending = None
        def flush():
            global first_paragraph
            if pending:
                text,style,size = pending
                paragraph = doc.add_paragraph(text,style)
                if first_paragraph and index:
                    paragraph.paragraph_format.page_break_before=True
                first_paragraph=False
                if size <= 7.5:
                    for run in paragraph.runs: run.font.size = Pt(9)
                all_text.append(text)
        previous = None
        for top, kind, value in events:
            if kind == 'image':
                flush();pending=None;previous=None
                location,data = value
                ratio = location['width']/location['height']
                height = min(4.0,16/ratio)
                paragraph=doc.add_paragraph()
                paragraph.paragraph_format.keep_with_next=True
                paragraph.add_run().add_picture(BytesIO(data),width=Cm(height*ratio),height=Cm(height))
                continue
            text,size = value['text'],value['size']
            style = 'Title' if size>23 else 'Heading 1' if size>15 else 'Heading 2' if value['bold'] and size>=9.5 else 'Normal'
            # Rejoin PDF soft wraps while keeping original paragraph boundaries.
            continuation = pending and previous and style=='Normal' and pending[1]=='Normal' and abs(size-previous['size'])<.2 and top-previous['bottom']<6 and not re.match(r'^(•|\d+[、.]|来源：|预览：)',text)
            if continuation:
                separator = ' ' if pending[0][-1:].isascii() and text[:1].isascii() else ''
                pending=(pending[0]+separator+text,style,size)
            else:
                flush();pending=(text,style,size)
            previous=value
        flush()
doc.core_properties.title = '图像生成任务库 Agent 批量执行指南'
doc.core_properties.subject = '2026-10-07 PDF 快照的可编辑版本'
doc.save(out)
assert sum('任务 ID：' in t for t in all_text)==18
assert len(doc.inline_shapes)==sum(len(p.images) for p in reader.pages)
print(f'{out}: 18 tasks, {len(doc.inline_shapes)} reference images, editable text')
