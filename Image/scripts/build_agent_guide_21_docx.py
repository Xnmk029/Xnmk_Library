"""Extend the retained guide; preserve its 18 tasks and embedded reference media."""
from pathlib import Path
import hashlib,json,re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from io import BytesIO

root=Path(__file__).resolve().parents[2]
reference=root/'output/docx/image-agent-batch-guide-20261007.docx'
out=root/'output/docx/image-agent-batch-guide-20261010-21tasks.docx'
doc=Document(reference)
new=[('VocaloidEnsemble','V家六位虚拟歌手同屏'),('GenshinEnsemble','原神六位角色同屏'),('ParisTrocaderoView','巴黎特罗卡德罗实景')]

def before(anchor,text,style='Normal'):
 return anchor.insert_paragraph_before(text,style)
def after(anchor,text,style='Normal'):
 p=doc.add_paragraph(text,style)
 anchor._p.addnext(p._p)
 return p

# Local edits to the cover and execution rules; retain the original task wording.
for p in doc.paragraphs:
 if p.text=='2026-10-07 · 工作区快照':p.text='2026-10-10 · 21题派发版'
 elif p.text.startswith('18 个题组'):
  p.text='21 个题组，覆盖写实生成、照片编辑、平面排版、漫画叙事、提示词与质量强度专题，以及世界知识。前18题保留2026-10-07指南内容；新增V家同屏、原神同屏和巴黎实景三题。循环精灵图、连续场景和旧渡船漫画不纳入本批次。'
 elif p.text.startswith(('PDF 中嵌入','PDF中嵌入','本文件已嵌入全部参考图片')):
  p.text='本DOCX内嵌原有题目的13处参考图，供识别和提取。若可访问题面中的工作区路径，优先读取原始文件；跨电脑执行时，可将DOCX作为ZIP解包，从word/media中提取下方注明的图片文件。不要使用页面截图或预览缩略图代替内嵌图。保持各题参考图顺序，提取后记录实际输入尺寸与SHA-256。'
  after(p,'第19至21题是无参考图的世界知识测试。整个生成过程禁止搜索、浏览器、地图、街景和外部素材检索；仅提交本题提示词。文档中没有这三题的生成样例。每题只生成一次，保留首轮缺陷，不自动修图或重试。')
 if p.text=='四种测试关系':
  before(p,'世界知识','Heading 2')
  before(p,'V家同屏、原神同屏和巴黎实景。检查角色身份与外观知识、多人属性绑定，以及真实地标与空间关系。三题均禁止搜索且无参考图；知识正确性与画面质量分开评审。')

# Move the existing topic note into the comparison section to avoid a tiny
# continuation page after the expanded cover; its wording stays unchanged.
paragraphs=list(doc.paragraphs)
topic=next(p for p in paragraphs if p.text=='专题批次')
topic_body=paragraphs[paragraphs.index(topic)+1]
comparison=next(p for p in doc.paragraphs if p.text=='四种测试关系')
comparison._p.addprevious(topic._p)
comparison._p.addprevious(topic_body._p)
in_overview=False
for p in doc.paragraphs:
 if p.text=='测试部分与对比方式':in_overview=True;continue
 if p.text=='题目索引':in_overview=False
 if in_overview:
  p.paragraph_format.space_before=Pt(3 if p.style.name=='Heading 2' else 0)
  p.paragraph_format.space_after=Pt(3)

# Update index using its existing body pattern.
anchor=next(p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('01 '))
for i,(tid,title) in enumerate(new,19):
 before(anchor,f'{i:02d}  {title}  |  {tid}')
 p=before(anchor,'可执行：文生图，禁止搜索，无参考图')
 for r in p.runs:r.font.size=Pt(9)

# Add media locators adjacent to each retained drawing, without resampling media.
media=[]
for p in list(doc.paragraphs):
 for blip in p._p.xpath('.//a:blip'):
  rid=blip.get(qn('r:embed'))
  part=doc.part.related_parts[rid]
  with Image.open(BytesIO(part.blob)) as im:w,h=im.size
  name=str(part.partname).lstrip('/')
  label=f'内嵌资产 {len(media)+1:02d}：{name}；{w}×{h}；保持原图顺序。'
  pp=before(p,label)
  for r in pp.runs:r.font.size=Pt(9)
  pp.paragraph_format.keep_with_next=True
  media.append(dict(file=name,width=w,height=h,sha256=hashlib.sha256(part.blob).hexdigest()))
assert len(media)==13

# Insert new task pages, reusing the retained heading/body styles.
for i,(tid,title) in enumerate(new,19):
 p=doc.add_paragraph(f'{i:02d}  {title}','Heading 1');p.paragraph_format.page_break_before=True
 p=doc.add_paragraph(f'任务 ID：{tid} | 版本：1.0 | 分类：世界知识 文生图')
 for r in p.runs:r.font.size=Pt(9)
 doc.add_paragraph('输入与执行条件','Heading 2')
 doc.add_paragraph('无参考图。禁止搜索或检索任何外部资料；禁止把其他模型的生成结果或评审资料作为本题输入。使用下列提示词原文，每个授权模型配置输出一张。')
 doc.add_paragraph('提示词原文','Heading 2')
 prompt=(root/'Image'/tid/'PROJECT_PROMPT.md').read_text(encoding='utf-8').strip()
 for block in prompt.split('\n\n'):
  if block.startswith('# '):continue
  doc.add_paragraph(block)
 doc.add_paragraph('交付与记录','Heading 2')
 doc.add_paragraph('保存原始PNG、实际调用的完整提示词和meta.json。记录请求与返回模型ID、质量、实际尺寸、耗时、工具轨迹及输出SHA-256。若工具未公开模型或内部搜索轨迹，明确标为未公开或未验证；不要推定。知识正确性与画面质量分开评审。')

# Keep page geometry and the source's type hierarchy; remove inherited title rules.
for style in doc.styles:
 for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
for border in doc.element.xpath('//w:pBdr'):border.getparent().remove(border)
for section in doc.sections:
 for p in section.footer.paragraphs:
  for r in p.runs:r.text=r.text.replace('2026-10-07','2026-10-10')
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.subject='21题派发版，内嵌13处原始参考图，无生成样例'
doc.save(out)

text='\n'.join(p.text for p in doc.paragraphs)
assert text.count('任务 ID：')==21
assert len(doc.inline_shapes)==13
assert all(t in text for t,_ in new)
# The output keeps every image binary from the retained companion document.
import zipfile
with zipfile.ZipFile(reference) as a,zipfile.ZipFile(out) as b:
 for name in a.namelist():
  if name.startswith('word/media/'):assert a.read(name)==b.read(name)
audit=root/'tmp/docx-guide-21-qa';audit.mkdir(parents=True,exist_ok=True)
(audit/'media-manifest.json').write_text(json.dumps(media,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'{out}: 21 tasks, 13 reference placements, image binaries unchanged')
